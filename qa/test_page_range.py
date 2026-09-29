#!/usr/bin/env python3
"""Тест выборочной печати страниц (print_engine.sheet_to_pdf).

Бланк на несколько страниц печатается в PDF целиком и по диапазонам:
счёт страниц в готовом PDF (pdfminer) должен совпадать с заказанным.

Запуск (песочница):
    LD_LIBRARY_PATH=/tmp/stubs QT_QPA_PLATFORM=offscreen \\
    OVERTIMETAB_SANDBOX_FONTS=1 python3 qa/test_page_range.py
"""
import os
import sys
import tempfile

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ.setdefault("OVERTIMETAB_SANDBOX_FONTS", "1")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from PySide6.QtGui import QGuiApplication

app = QGuiApplication([])

import openpyxl

import print_engine as pe

TEMPLATE = os.path.join(ROOT, "Template.xlsx")


def multi_page_sheet(n_employees=260):
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb["Лист1"] if "Лист1" in wb.sheetnames else wb.active
    row = 16
    for i in range(n_employees):
        ws.cell(row, 3, i + 1)
        ws.cell(row, 4, "Тестов Тест Тестович %02d" % (i + 1))
        row += 1
    # печатная область шаблона зафиксирована до 19-й строки — растягиваем
    ws.print_area = "A1:AU%d" % (row - 1)
    return ws


def pdf_pages(path):
    from pdfminer.high_level import extract_pages
    return sum(1 for _ in extract_pages(path))


def main() -> int:
    ws = multi_page_sheet()
    tmp = tempfile.mkdtemp(prefix="page_range_")
    out = os.path.join(tmp, "t.pdf")

    total = pe.sheet_to_pdf(ws, out)
    assert total >= 4, "ожидали многостраничный бланк, вышло %d" % total
    assert pdf_pages(out) == total, "PDF-страниц (%d) != отчёт движка (%d)" % (
        pdf_pages(out), total)
    print("целиком: %d стр. ✓" % total)

    cases = [(1, 1, 1), (2, 3, 2), (total, total, 1), (total - 1, total, 2)]
    for pf, pt, want in cases:
        n = pe.sheet_to_pdf(ws, out, page_from=pf, page_to=pt)
        got = pdf_pages(out)
        assert n == want, "движок вернул %d (ожидали %d) для %d-%d" % (n, want, pf, pt)
        assert got == want, "в PDF %d страниц (ожидали %d) для %d-%d" % (
            got, want, pf, pt)
        print("диапазон %d-%d: %d стр. ✓" % (pf, pt, got))

    # за пределами бланка — пустой документ, без падения
    n = pe.sheet_to_pdf(ws, out, page_from=total + 5, page_to=total + 6)
    assert n == 0, "за пределами бланка ждали 0, вышло %d" % n
    print("за пределами бланка: 0 стр., без падения ✓")

    print("═══ ВЫБОРОЧНАЯ ПЕЧАТЬ СТРАНИЦ: ВСЁ ТОЧНО ═══")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
