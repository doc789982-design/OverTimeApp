#!/usr/bin/env python3
"""Дымовой тест печати: шаблон → заполненные строки → PDF и «принтер».

Проверяет главный каркас print_engine без базы данных:
  1. PDF-путь (QPdfWriter) рендерит страницы, сноски о статусах на месте;
  2. путь на реальный принтер при недоступном устройстве даёт понятную
     ошибку (RuntimeError), а не «тихую пустую печать» и не вылет.

Запуск (песочница):
    LD_LIBRARY_PATH=/tmp/stubs QT_QPA_PLATFORM=offscreen \\
    OVERTIMETAB_SANDBOX_FONTS=1 python3 qa/smoke_print.py
"""
import os
import sys
import tempfile

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ.setdefault("OVERTIMETAB_SANDBOX_FONTS", "1")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from PySide6.QtGui import QGuiApplication

app = QGuiApplication([])

import openpyxl

import print_engine as pe

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(ROOT, "Template.xlsx")


def filled_sheet():
    """Шаблон с тремя сотрудниками и статусами К/Б/О в днях."""
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb["Лист1"] if "Лист1" in wb.sheetnames else wb.active
    for i, r in enumerate((16, 17, 18), start=1):
        ws.cell(r, 3, i)                      # номер в первой колонке = строка данных
        ws.cell(r, 4, "Тестов Тест Тестович %d" % i)
    ws.cell(16, 10, "К")                      # командировка
    ws.cell(17, 12, "Б")                      # больничный
    ws.cell(18, 14, "О")                      # отпуск
    return ws


def main() -> int:
    ws = filled_sheet()
    out_pdf = os.path.join(tempfile.gettempdir(), "smoke_print.pdf")
    n = pe.sheet_to_pdf(ws, out_pdf, "landscape", "A4")
    assert n >= 1, "PDF: не собрались страницы"
    print("PDF: %d стр., %d КБ" % (n, os.path.getsize(out_pdf) // 1024))

    raw = open(out_pdf, "rb").read()
    assert raw[:5] == b"%PDF-", "PDF: битый заголовок"

    # сноска о статусах должна попасть в PDF (проверяем текстом)
    try:
        from pdfminer.high_level import extract_pages
        from pdfminer.layout import LTTextContainer
        texts = []
        for page in extract_pages(out_pdf):
            for el in page:
                if isinstance(el, LTTextContainer):
                    texts.append(" ".join(el.get_text().split()))
        blob = " ".join(texts)
        assert "К - командировка" in blob, "PDF: нет сноски «К - командировка»"
        assert "О - отпуск" in blob, "PDF: нет сноски «О - отпуск»"
        print("PDF: сноски о статусах на месте")
    except ImportError:
        print("PDF: pdfminer недоступен, текст сносок не проверяли")

    # «Реальный принтер», которого нет: обязана быть понятная ошибка,
    # а не тихая пустая печать (и не вылет)
    try:
        pe.print_sheet_to_printer(ws, "Несуществующий принтер", 1, "", "",
                                  "landscape", "A4", True)
        raise AssertionError("принтер: ошибку не подняли (тихая печать?)")
    except RuntimeError as e:
        print("принтер недоступен → RuntimeError: %s" % e)
    print("═══ ДЫМОВОЙ ТЕСТ ПЕЧАТИ: ОК ═══")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
