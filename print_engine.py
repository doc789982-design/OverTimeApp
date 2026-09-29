# -*- coding: utf-8 -*-
"""
Собственный движок печати бланка — без Excel.

Логика: заполненный шаблон (openpyxl, как раньше) отрисовывается напрямую
на QPrinter через QPainter. Вся геометрия — ширины колонок, высоты строк,
объединённые ячейки, границы, заливки, шрифты, выравнивания, поворот
текста — читается из самого xlsx, поэтому бланк печатается таким, каким
его видит Excel, без запуска Excel.

Возможности:
  - печать на выбранный принтер (копии, диапазон страниц, дублет);
  - разбивка на страницы по высоте, как в Excel;
  - выгрузка в PDF и картинка (превью/тесты) тем же кодом.
"""
from __future__ import annotations

import datetime
import math
import os
import re
import tempfile
import uuid
from pathlib import Path

from utils import import_openpyxl

openpyxl = import_openpyxl()   # самолечение numpy-конфликта
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange

from PySide6.QtCore import QMarginsF, QPointF, QRectF
from PySide6.QtGui import QColor, QFont, QFontMetrics, QImage, QPainter, QPen, QTransform
from PySide6.QtPrintSupport import QPrinter

MDW = 7            # ширина цифры шрифта Normal (Calibri 11) в px при 96 dpi
DEFAULT_COL_CHARS = 8.43
PX_PER_PT = 96.0 / 72.0

# подмена шрифтов: приоритетные кандидаты (первый найденный в системе)
FONT_CHAINS = {
    "PT Astra Serif": ["PT Astra Serif", "PT Astra Serif Ext", "Times New Roman", "Liberation Serif"],
    "Calibri": ["Calibri", "Carlito", "Segoe UI", "Liberation Sans"],
}

_DAY_FORMULA = re.compile(r'^=IFERROR\(DAY\(([A-Z]{1,3})(\d+)\),')


def _col_chars_to_px(chars: float) -> int:
    return int(round(chars * MDW)) + 5


class SheetModel:
    """Геометрия и стили листа в «листовых пикселях» (96 dpi)."""

    def __init__(self, ws):
        self.ws = ws
        # ── область печати ──
        pa = ws.print_area
        if pa:
            s = (pa if isinstance(pa, str) else pa[0]).split(",")[0].split("!")[-1].replace("$", "")
            cr = CellRange(s)
            self.c0, self.c1 = int(cr.min_col), int(cr.max_col)
            self.r0, self.r1 = int(cr.min_row), int(cr.max_row)
        else:
            self.c0, self.r0 = 1, 1
            self.c1, self.r1 = ws.max_column, ws.max_row

        # ── колонки: диапазоны <col min max> ──
        widths = {}
        for dim in ws.column_dimensions.values():
            if dim.width is None:
                continue
            lo = int(dim.min or 0)
            hi = int(dim.max or lo)
            for c in range(lo, hi + 1):
                widths[c] = (dim.width, bool(dim.hidden))
        default_w = ws.sheet_format.defaultColWidth or DEFAULT_COL_CHARS
        self.col_w = {}
        self.hidden_cols = set()
        for c in range(self.c0, self.c1 + 1):
            if c in widths:
                w, hidden = widths[c]
            else:
                w, hidden = default_w, False
            if hidden:
                self.col_w[c] = 0
                self.hidden_cols.add(c)
            else:
                self.col_w[c] = _col_chars_to_px(w)

        # ── ряды ──
        default_h = ws.sheet_format.defaultRowHeight or 15.0
        self.row_h = {}
        for r in range(self.r0, self.r1 + 1):
            d = ws.row_dimensions.get(r)
            h = (d.height if d and d.height else default_h) * PX_PER_PT
            self.row_h[r] = max(1, int(round(h)))

        # ── координаты ──
        self.x = {self.c0 - 1: 0}
        for c in range(self.c0, self.c1 + 1):
            self.x[c] = self.x[c - 1] + self.col_w[c]

        self.y = {self.r0 - 1: 0}
        for r in range(self.r0, self.r1 + 1):
            self.y[r] = self.y[r - 1] + self.row_h[r]

        self.width = self.x[self.c1]
        self.height = self.y[self.r1]
        # ── мержи (только внутри области печати) ──
        self.merges = {}
        for mr in ws.merged_cells.ranges:
            if mr.min_col > self.c1 or mr.max_col < self.c0 or mr.min_row > self.r1 or mr.max_row < self.r0:
                continue
            self.merges[(mr.min_row, mr.min_col)] = (mr.min_row, mr.max_row, mr.min_col, mr.max_col)
        self.merge_of = {}
        for a, (r1, r2, c1, c2) in self.merges.items():
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    self.merge_of[(r, c)] = a

        # ── параметры страницы ──
        ps = ws.page_setup
        try:
            self.scale = (float(ps.scale) if ps.scale else 100.0) / 100.0
        except Exception:
            self.scale = 1.0
        m = ws.page_margins
        self.margins_mm = (25.4 * (m.left or 0.7), 25.4 * (m.right or 0.7),
                           25.4 * (m.top or 0.75), 25.4 * (m.bottom or 0.75))

        # последним: высоты строк без явной высоты (Excel auto-fit)
        self._autofit_default_rows()

    def _autofit_default_rows(self):
        """Excel подгоняет высоту строк БЕЗ явной высоты под содержимое
        (блок подписей после вставки строк сотрудников теряет высоты).
        Оцениваем по размеру шрифта: высота строки ≈ 1.35 × размер,
        символ ≈ 0.55 × размер; ячейки с переносом — по числу строк."""
        for r in range(self.r0, self.r1 + 1):
            d = self.ws.row_dimensions.get(r)
            if d is not None and d.height:
                continue  # явная высота — не трогаем
            need = 0.0
            for c in range(self.c0, self.c1 + 1):
                if c in self.hidden_cols or not self.is_anchor(r, c):
                    continue
                cell = self.ws.cell(r, c)
                v = self.value_for_print(cell)
                if not v:
                    continue
                size_px = max(6.0, (cell.font.sz or 11.0) * PX_PER_PT)
                line_h = size_px * 1.35
                if cell.alignment.wrap_text:
                    rect = self.cell_rect(r, c)
                    cpl = max(1, int(max(8.0, rect.width() - 3) / (size_px * 0.55)))
                    lines = 0
                    for para in str(v).split("\n"):
                        lines += max(1, math.ceil(len(para) / cpl))
                else:
                    # без переноса текст вытекает на пустых соседей ВБОК —
                    # высота строки от этого не растёт
                    lines = len(str(v).split("\n"))
                need = max(need, lines * line_h + 4)
            if need > self.row_h[r]:
                self.row_h[r] = int(round(need))

        # координаты Y изменились — пересобираем
        self.y = {self.r0 - 1: 0}
        for r in range(self.r0, self.r1 + 1):
            self.y[r] = self.y[r - 1] + self.row_h[r]
        self.height = self.y[self.r1]


    # ---------- строки данных и колонки компенсаций ----------

    def _data_rows(self):
        """Строки сотрудников: в первой колонке области печати стоит номер."""
        rows = []
        for r in range(self.r0, self.r1 + 1):
            v = self.ws.cell(r, self.c0).value
            if isinstance(v, int) or (isinstance(v, str) and v.strip().isdigit()):
                rows.append(r)
        return rows

    def first_data_row(self):
        rows = self._data_rows()
        return rows[0] if rows else 0

    def last_data_row(self):
        rows = self._data_rows()
        return rows[-1] if rows else 0

    def compensation_trios(self):
        """Тройки колонок (сверх / ночные / дни) групп «на начало» и «на конец».

        Группы находятся по широким титулам «Количество подлежащих
        компенсации часов (дней)…» (мерж на 3 колонки выше строк данных):
        левая группа — «на начало месяца», правая — «на конец месяца».
        Подколонки определяются по ключевым словам подзаголовков
        (вертикальные заголовки «на конец» — такие же ячейки).
        """
        first = self.first_data_row()
        if not first:
            return {}
        groups = []
        for a, (r1, r2, c1, c2) in self.merges.items():
            if r2 >= first or c2 - c1 + 1 != 3:
                continue
            v = self.ws.cell(r1, c1).value
            if isinstance(v, str) and "подлежащих компенсации" in v:
                groups.append((c1, c2, r2))
        if not groups:
            return {}
        groups.sort()

        def trio_for(c1, c2, title_bottom):
            trio = {}
            for r in range(title_bottom + 1, first):
                for c in range(c1, c2 + 1):
                    if c in self.hidden_cols:
                        continue
                    v = self.ws.cell(r, c).value
                    if not isinstance(v, str):
                        continue
                    low = v.lower()
                    for kw, key in (("сверх", "ot"), ("ночн", "hours"), ("нерабоч", "days")):
                        if kw in low and key not in trio:
                            trio[key] = c
            return trio if len(trio) == 3 else None

        res = {}
        for label, g in zip(("start", "end"), (groups[0], groups[-1])):
            t = trio_for(*g)
            if t:
                res[label] = t
        return res

    STATUS_LETTERS = ("К", "Б", "О", "В")
    STATUS_WORDS = {"ОТПУСК": "О", "БОЛЬНИЧНЫЙ": "Б", "КОМАНДИРОВКА": "К"}

    def letters_in_rows(self, rows):
        """Буквы статусов (К/Б/О/В), встречающиеся в данных строках.

        Учитываются и одиночные буквы в ячейках дней, и слова полос
        («ОТПУСК» и т.д. — та же буква).
        """
        found = set()
        for r in rows:
            for c in range(self.c0, self.c1 + 1):
                if c in self.hidden_cols:
                    continue
                v = self.value_for_print(self.ws.cell(r, c))
                if not isinstance(v, str) or not v:
                    continue
                for ln in v.split("\n"):
                    if ln.strip() in self.STATUS_LETTERS:
                        found.add(ln.strip())
                for word, letter in self.STATUS_WORDS.items():
                    if word in v:
                        found.add(letter)
        return found

    # ---------- содержимое ----------

    def cell_rect(self, r, c):
        a = self.merge_of.get((r, c))
        if a:
            r1, r2, c1, c2 = self.merges[a]
            return QRectF(self.x[c1 - 1], self.y[r1 - 1],
                          self.x[c2] - self.x[c1 - 1], self.y[r2] - self.y[r1 - 1])
        return QRectF(self.x[c - 1], self.y[r - 1], self.col_w[c], self.row_h[r])

    def is_anchor(self, r, c):
        a = self.merge_of.get((r, c))
        return a is None or a == (r, c)

    def value_for_print(self, cell):
        """Значение ячейки, как его показывает Excel на печати."""
        nf = (cell.number_format or "").strip()
        if nf == ";;;":          # формат «скрыть»
            return None
        v = cell.value
        if v is None:
            return None
        if isinstance(v, str):
            if v.startswith("="):
                # формулы шаблона скрыты форматом ;;; и не вычисляются;
                # на всякий случай поддержан единственный встречающийся вид
                m = _DAY_FORMULA.match(v)
                if m and nf != ";;;":
                    ref = self.ws[f"{m.group(1)}{m.group(2)}"].value
                    if isinstance(ref, (datetime.datetime, datetime.date)):
                        return str(ref.day)
                return None
            return v
        if isinstance(v, (datetime.datetime, datetime.date)):
            low = nf.lower()
            if ("d" in low and "m" not in low and "y" not in low) or low in ("general", "0"):
                return str(v.day)
            return v.strftime("%d.%m.%Y")
        if isinstance(v, float) and v == int(v):
            return str(int(v))
        return str(v)

    @staticmethod
    def color_of(cell, default=None):
        try:
            fill = cell.fill
            if fill and fill.fill_type == "solid":
                rgb = fill.start_color.rgb
                if isinstance(rgb, str) and len(rgb) >= 6:
                    return QColor("#" + rgb[-6:]), True
                th = getattr(fill.start_color, "theme", None)
                if th == 0:  # background1 — белый лист
                    return QColor("#ffffff"), True
        except Exception:
            pass
        return (default, False)


class SheetRenderer:
    """Рисует SheetModel на QPainter в листовых пикселях."""

    PEN_HAIR = 0.33    # 0.25 pt в листовых px
    PEN_THIN = 1.0     # 0.75 pt

    def __init__(self, model: SheetModel, font_hook=None):
        self.m = model
        self.font_hook = font_hook or (lambda name: name)   # для тестов: подмена шрифта
        self._fnt = {}
        # {(row, "start"/"end"): "N дн."} — тройка колонок группы разделена
        # ОДНОЙ наклонной линией: сверху — текущие значения трёх граф,
        # снизу — одно число: часы и дни, переведённые в дни (как «Всего дней»)
        self.days_overlay = {}
        self._trios = None

    def qfont(self, cell):
        f = cell.font
        name = f.name or "Calibri"
        chain = FONT_CHAINS.get(name)
        if chain:
            for cand in chain:
                fname = self.font_hook(cand)
                if fname:
                    name = fname
                    break
            else:
                name = self.font_hook(chain[-1]) or name
        else:
            name = self.font_hook(name) or name
        size = max(6, int(round((f.sz or 11.0) * PX_PER_PT)))
        key = (name, size, bool(f.b))
        if key not in self._fnt:
            font = QFont(name)
            font.setPixelSize(size)
            font.setBold(bool(f.b))
            font.setItalic(bool(f.i))
            self._fnt[key] = font
        return self._fnt[key]

    # ---------- страницы ----------

    def page_rows(self, page_h_px):
        """Разбивает видимые строки print_area на страницы по высоте."""
        pages, cur, cur_h = [], [], 0.0
        for r in range(self.m.r0, self.m.r1 + 1):
            h = self.m.row_h[r]
            if cur and cur_h + h > page_h_px:
                pages.append(cur)
                cur, cur_h = [], 0.0
            cur.append(r)
            cur_h += h
        if cur:
            pages.append(cur)

        # Подписант не должен оставаться один на последнем листе: если туда
        # не попало ни одной строки сотрудника, последний сотрудник (и всё,
        # что шло за ним на предыдущем листе) переезжает вместе с ним.
        last_emp = self.m.last_data_row()
        if last_emp and len(pages) >= 2 and pages[-1][0] > last_emp:
            prev = pages[-2]
            if last_emp in prev:
                i = prev.index(last_emp)
                pages[-2] = prev[:i]
                pages[-1] = prev[i:] + pages[-1]
        return pages

    # ---------- отрисовка ----------

    def paint(self, painter: QPainter, rows, x_shift=0.0):
        m = self.m
        painter.save()
        painter.translate(-x_shift, -m.y[rows[0] - 1])

        # 1) заливки
        for r in rows:
            for c in range(m.c0, m.c1 + 1):
                if c in m.hidden_cols:
                    continue
                cell = m.ws.cell(r, c)
                color, ok = m.color_of(cell)
                if ok and color is not None:
                    rect = m.cell_rect(r, c)
                    painter.fillRect(rect, color)

        # 2) границы
        for r in rows:
            for c in range(m.c0, m.c1 + 1):
                if c in m.hidden_cols:
                    continue
                self._borders(painter, r, c)

        # 3) текст
        for r in rows:
            for c in range(m.c0, m.c1 + 1):
                if c in m.hidden_cols or not m.is_anchor(r, c):
                    continue
                cell = m.ws.cell(r, c)
                text = m.value_for_print(cell)
                if text is None or text == "":
                    continue
                self._text(painter, r, c, cell, str(text))

        # тройки «на начало»/«на конец»: одна линия на тройку + дни под ней
        if self.days_overlay:
            rows_set = set(rows)
            trios = self._trio_groups()
            for (row, group), days in self.days_overlay.items():
                if row not in rows_set:
                    continue
                trio = trios.get(group)
                if trio:
                    self._paint_trio_days(painter, row, group, trio, days)
        painter.restore()

    def _draw_rect(self, r, c):
        """Прямоугольник для текста: без переноса Excel позволяет тексту
        выезжать на ПУСТЫЕ соседние ячейки (пока не встретится занятая)."""
        m = self.m
        rect = m.cell_rect(r, c)
        cell = m.ws.cell(r, c)
        al = cell.alignment
        if al.wrap_text:
            return rect
        if al.horizontal == "right":
            cc = c - 1
            while cc >= m.c0:
                if m.ws.cell(r, cc).value is not None:
                    break  # скрытые колонки текст перетекает, как в Excel
                cc -= 1
            if cc < c - 1:
                left = m.x[cc] if cc >= m.c0 else 0
                if left < rect.left():
                    rect = QRectF(left, rect.top(), rect.right() - left, rect.height())
        else:
            cc = c + 1
            while cc <= m.c1:
                if m.ws.cell(r, cc).value is not None:
                    break  # скрытые колонки текст перетекает, как в Excel
                cc += 1
            right = m.x[cc - 1]
            if right > rect.right():
                rect = QRectF(rect.left(), rect.top(), right - rect.left(), rect.height())
        return rect

    FOOTNOTE_LABELS = {"К": "К - командировка", "Б": "Б - больничный",
                       "О": "О - отпуск", "В": "В - выходной"}

    def footnote_font(self, dpi=96):
        """Шрифт сносок: семейство из ячеек данных, чуть мельче.
        dpi — разрешение устройства (шрифт масштабируется под него)."""
        f = self.m.first_data_row()
        cell = self.m.ws.cell(f, self.m.c0) if f else None
        base = self.qfont(cell) if cell is not None else QFont("Calibri")
        font = QFont(base)
        font.setPixelSize(max(6, int(base.pixelSize() * 0.8 * dpi / 96.0)))
        font.setBold(False)
        return font

    def paint_footnote(self, painter, font, baseline_y, letters):
        """Сноска о статусах одной строкой внизу листа."""
        text = self.footnote_text(letters)
        painter.setFont(font)
        painter.setPen(QColor("#333333"))
        painter.drawText(0, int(round(baseline_y)), text)

    def footnote_text(self, letters):
        """«К - командировка   Б - больничный …» в фиксированном порядке."""
        parts = [self.FOOTNOTE_LABELS[l] for l in self.STATUS_LETTERS_ORDER if l in letters]
        return "    ".join(parts)

    STATUS_LETTERS_ORDER = ("К", "Б", "О", "В")

    def _trio_groups(self):
        if self._trios is None:
            try:
                self._trios = self.m.compensation_trios()
            except Exception:
                self._trios = {}
        return self._trios

    def _trio_of_cell(self, r, c):
        """(группа, тройка), если ячейка входит в тройку с оверлеем этой строки."""
        for group, trio in self._trio_groups().items():
            if c in (trio["ot"], trio["hours"], trio["days"]):
                if (r, group) in (self.days_overlay or {}):
                    return group, trio
        return None, None

    def _trio_rect(self, r, trio):
        """Общий прямоугольник трёх колонок тройки в строке r."""
        m = self.m
        cols = sorted((trio["ot"], trio["hours"], trio["days"]))
        return QRectF(m.x[cols[0] - 1], m.y[r - 1],
                      m.x[cols[-1]] - m.x[cols[0] - 1], m.row_h[r])

    @staticmethod
    def _trio_line_y(rect, x):
        """Y наклонной линии (~10°, слева ниже — справа выше) в точке x."""
        mid = (rect.top() + rect.bottom()) / 2
        dy = math.tan(math.radians(10.0)) * rect.width() / 2.0
        t = (x - rect.left()) / rect.width()
        return mid + dy - t * 2 * dy

    def _trio_edge_clips(self, r):
        """{x внутренней границы тройки: y линии} — под линией границ нет."""
        if not self.days_overlay:
            return {}
        cached = getattr(self, "_clips_cache", None)
        if cached is None:
            cached = self._clips_cache = {}
        if r in cached:
            return cached[r]
        res = {}
        for group, trio in self._trio_groups().items():
            if (r, group) not in self.days_overlay:
                continue
            rect = self._trio_rect(r, trio)
            cols = sorted((trio["ot"], trio["hours"], trio["days"]))
            for cc in cols[1:]:
                x = self.m.x[cc - 1]
                res[x] = self._trio_line_y(rect, x)
        cached[r] = res
        return res

    def _text_top_half(self, painter, r, c, cell, text, fm, trio):
        """Значение ячейки тройки — в верхней половине (над линией)."""
        rect = self.m.cell_rect(r, c)
        al = cell.alignment
        line_y = self._trio_line_y(self._trio_rect(r, trio), rect.center().x())
        box = QRectF(rect.left(), rect.top() + 1,
                     rect.width(), max(8.0, line_y - rect.top() - 4))
        lines = (self._wrap_lines(text, fm, box.width() - 3)
                 if al.wrap_text else text.split("\n"))
        total = len(lines) * fm.height()
        y = box.top() + max(0.0, (box.height() - total) / 2)
        painter.setPen(QColor("#111111"))
        for ln in lines:
            tw = fm.horizontalAdvance(ln)
            if al.horizontal == "center":
                x = box.left() + (box.width() - tw) / 2
            elif al.horizontal == "right":
                x = box.right() - tw
            else:
                x = box.left()
            painter.drawText(int(x), int(y + fm.ascent()), ln)
            y += fm.height()

    def _paint_trio_days(self, painter, row, group, trio, days):
        """Одна наклонная линия через все три колонки + одно число дней под ней."""
        rect = self._trio_rect(row, trio)
        mid = (rect.top() + rect.bottom()) / 2
        dy = math.tan(math.radians(10.0)) * rect.width() / 2.0
        painter.setPen(QPen(QColor("#808080"), 0.75))
        painter.drawLine(QPointF(rect.left(), mid + dy), QPointF(rect.right(), mid - dy))

        cell = self.m.ws.cell(row, trio["ot"])
        font = QFont(self.qfont(cell))   # копия — не портим кэш шрифтов
        font.setBold(True)
        fm = QFontMetrics(font)
        painter.setFont(font)
        painter.setPen(QColor("#111111"))
        bot_top = mid + dy + 2
        bot_h = rect.bottom() - bot_top - 2
        tw = fm.horizontalAdvance(days)
        x = rect.left() + (rect.width() - tw) / 2
        y = bot_top + max(0.0, (bot_h - fm.height()) / 2)
        painter.drawText(int(x), int(y + fm.ascent()), days)

    @staticmethod
    def _char_split(word, fm, maxw):
        """Режет слово шире ячейки по символам (Excel делает так же)."""
        out, cur = [], ""
        for ch in word:
            if cur and fm.horizontalAdvance(cur + ch) > maxw:
                out.append(cur)
                cur = ch
            else:
                cur += ch
        if cur:
            out.append(cur)
        return out

    def _wrap_lines(self, text, fm, maxw):
        """Перенос по словам; неумещающееся слово — по символам."""
        lines = []
        for para in text.split("\n"):
            cur = ""
            for word in para.split(" "):
                cand = (cur + " " + word).strip()
                if fm.horizontalAdvance(word) > maxw:
                    # слово само по себе шире ячейки
                    if cur:
                        lines.append(cur)
                        cur = ""
                    lines.extend(self._char_split(word, fm, maxw))
                elif fm.horizontalAdvance(cand) <= maxw or not cur:
                    cur = cand
                else:
                    lines.append(cur)
                    cur = word
            if cur:
                lines.append(cur)
        return lines

    def _borders(self, painter, r, c):
        m = self.m
        cell = m.ws.cell(r, c)
        a = m.merge_of.get((r, c))
        mr = m.merges.get(a) if a else None
        rect = m.cell_rect(r, c)

        def pen(side):
            b = getattr(cell.border, side, None)
            if not b or not b.style:
                return None
            color = "#808080" if b.style == "hair" else "#1a1a1a"
            try:
                if b.color and isinstance(b.color.rgb, str) and len(b.color.rgb) >= 6:
                    color = "#" + b.color.rgb[-6:]
            except Exception:
                pass
            width = self.PEN_HAIR if b.style == "hair" else self.PEN_THIN
            return QPen(QColor(color), width)

        def draw(side, x1, y1, x2, y2):
            p = pen(side)
            if p is None:
                return
            painter.setPen(p)
            painter.drawLine(x1, y1, x2, y2)

        def draw_v(side, x, y1, y2):
            """Вертикальная граница; внутренние границы троек обрезаются
            по наклонной линии — под ней тройка становится одним полем."""
            p = pen(side)
            if p is None:
                return
            clip = self._trio_edge_clips(r).get(x)
            if clip is not None and y2 > clip:
                y2 = clip
            if y2 <= y1:
                return
            painter.setPen(p)
            painter.drawLine(x, y1, x, y2)

        left_c = mr[2] if mr else c
        right_c = mr[3] if mr else c
        top_r = mr[0] if mr else r
        bot_r = mr[1] if mr else r
        draw_v("left", rect.left(), rect.top(), rect.bottom())
        if c == right_c:
            draw_v("right", rect.right(), rect.top(), rect.bottom())
        draw("top", rect.left(), rect.top(), rect.right(), rect.top())
        if r == bot_r:
            draw("bottom", rect.left(), rect.bottom(), rect.right(), rect.bottom())

    def _text(self, painter, r, c, cell, text):
        m = self.m
        rect = self._draw_rect(r, c)   # с вытеканием на пустых соседей
        inset = 1.5
        box = rect.adjusted(inset, inset, -inset, -inset)
        font = self.qfont(cell)
        al = cell.alignment
        painter.setFont(font)
        painter.setPen(QColor("#111111"))
        fm = QFontMetrics(font)
        painter.save()
        # Всё, что не влезло, обрезается границей ячейки — как в Excel
        painter.setClipRect(rect)

        group, trio = self._trio_of_cell(r, c)
        if group is not None:
            # ячейка тройки с оверлеем: значение — в верхней половине,
            # линию и дни под ней рисует paint() один раз на тройку
            self._text_top_half(painter, r, c, cell, text, fm, trio)
            painter.restore()
            return

        rot = al.textRotation or 0
        if rot == 90:
            # Вертикальный текст переносится по ВЫСОТЕ ячейки
            if al.wrap_text:
                lines = self._wrap_lines(text, fm, box.height())
            else:
                lines = text.split("\n")
            painter.translate(rect.center())
            painter.rotate(-90)
            w, h = box.height(), box.width()
            total = len(lines) * fm.height()
            y = -total / 2
            for ln in lines:
                tw = fm.horizontalAdvance(ln)
                x = -tw / 2 if al.horizontal in (None, "center") else (-w / 2 if al.horizontal == "left" else w / 2 - tw)
                painter.drawText(int(x), int(y + fm.ascent()), ln)
                y += fm.height()
            painter.restore()
            return

        # Перенос по словам; слово шире ячейки режется по символам —
        # как в Excel; иначе текст выезжает на соседние клетки
        if al.wrap_text:
            lines = self._wrap_lines(text, fm, box.width())
        else:
            lines = text.split("\n")

        total = len(lines) * fm.height()
        # Умолчание Excel: без явного выравнивания текст прижат К НИЗУ ячейки
        if al.vertical == "center":
            y = box.top() + (box.height() - total) / 2
        elif al.vertical == "top":
            y = box.top()
        else:
            y = box.bottom() - total
        for ln in lines:
            tw = fm.horizontalAdvance(ln)
            if al.horizontal == "center":
                x = box.left() + (box.width() - tw) / 2
            elif al.horizontal == "right":
                x = box.right() - tw
            else:
                x = box.left()
            painter.drawText(int(x), int(y + fm.ascent()), ln)
            y += fm.height()
        painter.restore()


# ──────────────────────────────────────────────────────────────────
# Вывод: принтер / PDF / картинка
# ──────────────────────────────────────────────────────────────────

def _apply_pages(pdevice, model, orientation, paper_size):
    from PySide6.QtGui import QPageLayout, QPageSize
    size = QPageSize(QPageSize.A3 if paper_size == "A3" else QPageSize.A4)
    orient = QPageLayout.Landscape if orientation == "landscape" else QPageLayout.Portrait
    pdevice.setPageSize(size)
    pdevice.setPageOrientation(orient)
    ml, mr, mt, mb = model.margins_mm
    pdevice.setPageMargins(QMarginsF(ml, mt, mr, mb), QPageLayout.Millimeter)


def render_to_device(model, pdevice, page_from=None, page_to=None, days_overlay=None):
    """Рисует лист на QPrinter/QPdfWriter. Возвращает число напечатанных страниц."""
    renderer = SheetRenderer(model)
    renderer.days_overlay = days_overlay or {}
    painter = QPainter(pdevice)
    if not painter.isActive():
        # Устройство не дало начать печать (нет движка/принтер недоступен) —
        # не молчим и не рисуем в пустоту, а уходим в понятную ошибку.
        raise RuntimeError("Устройство печати недоступно (не удалось начать печать)")
    try:
        # печатная область в листовых px с учётом масштаба шаблона
        pr = pdevice.pageLayout().paintRectPixels(pdevice.resolution())
        k_base = pdevice.resolution() / 96.0 * model.scale
        avail_w_px = pr.width() / k_base
        if model.width > avail_w_px:          # чуть не влезает — прижимаем ширину
            k_base *= avail_w_px / model.width
        page_h_px = pr.height() / k_base
        pages = renderer.page_rows(page_h_px)
        pages = pages[int(page_from) - 1: int(page_to)] if (page_from and page_to) else pages
        cm_px = pdevice.resolution() / 2.54          # пикселей устройства в сантиметре
        # ВАЖНО: у QPainter на устройствах печати начало координат — угол
        # ОБЛАСТИ КОНТЕНТА (после полей), а не угол листа. Считаем базовую
        # линию сноски от края листа и переводим в координаты художника.
        fr = pdevice.pageLayout().fullRectPixels(pdevice.resolution())
        fn_baseline = fr.bottom() - 0.5 * cm_px - pr.y()
        fn_font = renderer.footnote_font(pdevice.resolution())
        for i, rows in enumerate(pages):
            if i:
                pdevice.newPage()
            painter.save()
            painter.setTransform(QTransform())
            painter.scale(k_base, k_base)
            renderer.paint(painter, rows)
            painter.restore()

            # Сноски о статусах — в самом низу листа (0,5 см от края),
            # только если на этой странице есть эти буквы
            letters = model.letters_in_rows(rows)
            if letters:
                renderer.paint_footnote(painter, fn_font, fn_baseline, letters)
        return len(pages)
    finally:
        painter.end()


def sheet_to_pdf(ws, out_path, orientation="landscape", paper_size="A4", days_overlay=None):
    from PySide6.QtGui import QPdfWriter
    model = SheetModel(ws)
    pdf = QPdfWriter(out_path)
    pdf.setResolution(96)
    _apply_pages(pdf, model, orientation, paper_size)
    return render_to_device(model, pdf, days_overlay=days_overlay)


def sheet_to_image(ws, out_path, dpi=96):
    """Превью первого фрагмента листа (для тестов)."""
    model = SheetModel(ws)
    renderer = SheetRenderer(model, font_hook=_HOOK or _system_font_hook)
    img = QImage(model.width, min(model.height, 5000), QImage.Format_RGB32)
    img.fill(QColor("#ffffff"))
    p = QPainter(img)
    p.setTransform(QTransform())
    renderer.paint(p, list(range(model.r0, model.r1 + 1)))
    p.end()
    img.save(out_path)
    return model


# подмена шрифтов для песочницы (нет PT Astra Serif); в Windows не активна
def _sandbox_hook(name):
    try:
        import glob
        ttf = glob.glob("/usr/local/lib/python3.11/dist-packages/matplotlib/"
                        "mpl-data/fonts/ttf/DejaVuSerif*.ttf")
        if not ttf:
            return name if name != "PT Astra Serif" else None
        from PySide6.QtGui import QGuiApplication, QFontDatabase
        if not QGuiApplication.instance():
            return None
        if not getattr(_sandbox_hook, "_loaded", False):
            from PySide6.QtGui import QFontDatabase as FD
            for f in ttf:
                FD.addApplicationFont(f)
            sans = glob.glob("/usr/local/lib/python3.11/dist-packages/matplotlib/"
                             "mpl-data/fonts/ttf/DejaVuSans.ttf")
            for f in sans:
                FD.addApplicationFont(f)
            _sandbox_hook._loaded = True
        return "DejaVu Serif" if "Serif" in name or "Astra" in name else "DejaVu Sans"
    except Exception:
        return None


def _system_font_hook(name):
    """Обычный режим: имя проходит, если шрифт есть в системе."""
    try:
        from PySide6.QtGui import QFontDatabase
        return name if name in QFontDatabase.families() else None
    except Exception:
        return name


_HOOK = _sandbox_hook if os.environ.get("OVERTIMETAB_SANDBOX_FONTS") else None


def _match_printer_name(requested, available_names):
    """Подбирает имя принтера в написании, которое знает Qt.

    Диалог печати получает имена из win32print, Qt — из своих источников:
    регистр, пробелы и сетевые пути («\\\\сервер\\принтер») могут
    отличаться. Сверяем без учёта регистра и хвоста сетевого пути.
    Возвращает имя из available_names или пустую строку.
    """
    req = (requested or "").strip()
    if not req:
        return ""
    low = req.lower()
    for n in available_names:
        if n.strip().lower() == low:
            return n
    tail = low.rsplit("\\", 1)[-1].strip()
    if tail:
        for n in available_names:
            if n.strip().lower().rsplit("\\", 1)[-1].strip() == tail:
                return n
    return ""


def print_sheet_to_printer(ws, printer_name, copies, page_from, page_to,
                           orientation, paper_size, collate, days_overlay=None):
    printer = QPrinter(QPrinter.HighResolution)
    requested = (printer_name or "").strip()
    if requested:
        # Qt молча оставляет принтер ПО УМОЛЧАНИЮ, если имя не совпало
        # с его списком (регистр/пробелы/сетевой путь из win32print).
        # Подбираем правильное написание и проверяем после установки.
        from PySide6.QtPrintSupport import QPrinterInfo
        names = [pi.printerName() for pi in QPrinterInfo.availablePrinters()]
        if names:
            matched = _match_printer_name(requested, names)
            if not matched:
                raise RuntimeError(
                    "Принтер «%s» не найден в системе (доступны: %s)"
                    % (requested, ", ".join(names[:6])))
            printer.setPrinterName(matched)
            if printer.printerName().strip().lower() != matched.strip().lower():
                raise RuntimeError(
                    "Система не дала выбрать принтер «%s»" % requested)
        else:
            printer.setPrinterName(requested)
    printer.setCopyCount(int(copies or 1))
    # QPrinter::setCollate убрали из Qt начиная с 6.11 — применяем, когда есть
    if hasattr(printer, "setCollate"):
        printer.setCollate(bool(collate))
    model = SheetModel(ws)
    _apply_pages(printer, model, orientation, paper_size)
    pf, pt = None, None
    try:
        if str(page_from).strip() and str(page_to).strip():
            pf, pt = int(page_from), int(page_to)
            printer.setFromTo(pf, pt)
    except ValueError:
        pass
    if not printer.isValid():
        raise RuntimeError("Принтер «%s» недоступен" % (printer_name or "по умолчанию"))
    return render_to_device(model, printer, pf, pt, days_overlay)


# ──────────────────────────────────────────────────────────────────
# Высокий уровень: база → заполненный бланк → печать
# ──────────────────────────────────────────────────────────────────

def _days_word(n: int) -> str:
    """Склонение: 1 день / 2 дня / 5 дней (как в панели итогов)."""
    n = abs(int(n))
    if n % 100 in (11, 12, 13, 14):
        return "дней"
    if n % 10 == 1:
        return "день"
    if n % 10 in (2, 3, 4):
        return "дня"
    return "дней"


def _compute_days_overlay(db_path, year, month, model):
    """Число дней для троек «на начало» и «на конец» месяца.

    Для каждой строки сотрудника считает одно число на группу — по той же
    формуле, что «Всего дней» в панели итогов (logic.total_overtime_days):
    ночные + сверх (сверх зажата в ноль) делятся на 8-часовой день
    с округлением вниз, плюс дни; учитываются остатки прошлого года.
    """
    trios = model.compensation_trios()
    emp_rows = model._data_rows()
    if not trios or not emp_rows:
        return {}
    from database import DB
    from logic import compute_month_summary, total_overtime_days

    keys = {
        "start": ("start_overtime", "prev_o_start", "start_hours", "prev_h_start", "start_days", "prev_d_start"),
        "end": ("end_overtime", "prev_o_end", "end_hours", "prev_h_end", "end_days", "prev_d_end"),
    }
    db = DB(db_path)
    try:
        emps = db.list_employees_for_month(year, month, active_only=True, search="")
    except Exception:
        db.close()
        return {}
    try:
        overlay = {}
        for i, row in enumerate(emp_rows):
            if i >= len(emps):
                break
            summ = compute_month_summary(db, int(emps[i]["id"]), year, month)
            for group in ("start", "end"):
                if group not in trios:
                    continue
                ot, ot_p, h, h_p, d, d_p = (int(summ[k] or 0) for k in keys[group])
                n = total_overtime_days(h, h_p, ot, ot_p, d, d_p)
                n = max(0, n)  # отрицательный остаток дней показываем нулём
                overlay[(row, group)] = "Всего %d %s" % (n, _days_word(n))
        return overlay
    finally:
        db.close()


def _build_sheet(db_path, year, month, template_path, sheet_name="Лист1"):
    """Формирует заполненный бланк во временный xlsx; возвращает (лист, путь)."""
    from database import DB
    from export import TemplateExporter

    temp_db = DB(db_path)
    try:
        temp_dir = Path(tempfile.gettempdir())
        temp_xlsx = temp_dir / ("overtimetab_qtprint_%s.xlsx" % uuid.uuid4().hex[:6])
        TemplateExporter.export(db=temp_db, year=year, month=month,
                                template_path=template_path, out_path=str(temp_xlsx),
                                sheet_name=sheet_name)
    finally:
        temp_db.close()

    wb = openpyxl.load_workbook(temp_xlsx)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
    return ws, temp_xlsx


def _cleanup_sheet(temp_xlsx):
    try:
        os.remove(str(temp_xlsx))
    except OSError:
        pass


def print_report(db_path, year, month, template_path, printer_name, copies,
                 page_from, page_to, orientation, paper_size, collate,
                 sheet_name="Лист1") -> int:
    """Формирует бланк и печатает его без Excel. Возвращает число страниц."""
    ws, temp_xlsx = _build_sheet(db_path, year, month, template_path, sheet_name)
    try:
        model = SheetModel(ws)
        overlay = _compute_days_overlay(db_path, year, month, model)
        n = print_sheet_to_printer(ws, printer_name, copies, page_from, page_to,
                                   orientation, paper_size, collate, overlay)
        return n
    finally:
        _cleanup_sheet(temp_xlsx)


def print_report_pdf(db_path, year, month, template_path, out_path,
                     orientation="landscape", paper_size="A4",
                     sheet_name="Лист1") -> int:
    """Печать в PDF-файл собственным средством (QPdfWriter, без драйвера).

    Путь для виртуальных PDF-принтеров («Microsoft Print to PDF» и т.п.):
    документ сохраняется в выбранный файл, системный диалог драйвера не нужен.
    """
    ws, temp_xlsx = _build_sheet(db_path, year, month, template_path, sheet_name)
    try:
        model = SheetModel(ws)
        overlay = _compute_days_overlay(db_path, year, month, model)
        return sheet_to_pdf(ws, out_path, orientation, paper_size, overlay)
    finally:
        _cleanup_sheet(temp_xlsx)
